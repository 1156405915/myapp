import hashlib
import importlib.util
import io
import json
import os
import subprocess
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, Mock, patch

WORKER_PATH = Path(__file__).resolve().parents[1] / "resources" / "document-processing" / "worker.py"
SPEC = importlib.util.spec_from_file_location("document_processing_worker", WORKER_PATH)
worker = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(worker)

try:
    import pypdfium2
    from pypdf import PdfWriter
    from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject, NumberObject
    from PIL import Image, ImageDraw
    from docx import Document
    HAS_DOCUMENT_DEPS = True
except ImportError:
    HAS_DOCUMENT_DEPS = False


NORMAL_TEXT = "This is a normal English document with sufficient readable text and no need for optical recognition."


def make_pdf(path, contents):
    writer = PdfWriter()
    for content in contents:
        mixed = content == "mixed"
        if mixed:
            content = NORMAL_TEXT
        page = writer.add_blank_page(width=300, height=200)
        font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                 NameObject("/Subtype"): NameObject("/Type1"),
                                 NameObject("/BaseFont"): NameObject("/Helvetica")})
        page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): writer._add_object(font)})})
        stream = DecodedStreamObject()
        if content == "scan":
            stream.set_data(b"0 g 30 60 200 30 re f")
        elif content:
            stream.set_data(b"BT /F1 8 Tf 10 130 Td (" + content.encode("ascii") + b") Tj ET")
        else:
            stream.set_data(b"")
        if mixed:
            image = DecodedStreamObject()
            image.set_data(b"\0" * 12)
            image.update({NameObject("/Type"): NameObject("/XObject"),
                          NameObject("/Subtype"): NameObject("/Image"),
                          NameObject("/Width"): NumberObject(2), NameObject("/Height"): NumberObject(2),
                          NameObject("/ColorSpace"): NameObject("/DeviceRGB"),
                          NameObject("/BitsPerComponent"): NumberObject(8)})
            page["/Resources"][NameObject("/XObject")] = DictionaryObject({NameObject("/Scan"): writer._add_object(image)})
            stream.set_data(stream.get_data() + b" q 80 0 0 40 20 20 cm /Scan Do Q")
        page[NameObject("/Contents")] = writer._add_object(stream)
    with path.open("wb") as stream:
        writer.write(stream)


def make_docx(path, picture=False):
    document = Document()
    document.add_paragraph("短中文正文 123，不应 OCR。")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "42"
    document.sections[0].header.paragraphs[0].text = "HEADER preserved"
    document.sections[0].footer.paragraphs[0].text = "FOOTER preserved"
    if picture:
        image = Image.new("RGB", (200, 80), "white")
        ImageDraw.Draw(image).rectangle((20, 20, 170, 50), fill="black")
        blob = io.BytesIO()
        image.save(blob, format="PNG")
        blob.seek(0)
        document.add_picture(blob)
    document.save(path)


@unittest.skipUnless(HAS_DOCUMENT_DEPS, "Requires local pypdfium2, pypdf, Pillow and python-docx")
class WorkerTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.input = self.root / "input.pdf"
        self.output = self.root / "output"
        self.ocr = Mock(return_value=[[[[0, 0], [100, 0], [100, 20], [0, 20]], "识别内容 123", 0.98]])

    def run_worker(self, name="example.pdf", source=None, ocr=None, emit=None, source_hash=None):
        source = source or self.input
        return worker.run(str(source), str(self.output), source_hash or hashlib.sha256(source.read_bytes()).hexdigest(),
                          name, ocr=self.ocr if ocr is None else ocr, emit=emit)

    def report_text(self):
        return (self.output / "full.txt").read_text(encoding="utf-8")

    def test_normal_english_does_not_initialize_ocr(self):
        make_pdf(self.input, [NORMAL_TEXT])
        result = self.run_worker()
        self.assertEqual(result["status"], "ready")
        self.assertEqual(result["pages"][0]["method"], "text")
        self.assertEqual(result["ocrPages"], 0)
        self.ocr.assert_not_called()

    def test_blank_page_is_visually_verified(self):
        make_pdf(self.input, [None])
        result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "blank")
        self.assertEqual(result["status"], "ready")
        self.ocr.assert_not_called()

    def test_chinese_name_numeric_text_requires_ocr_comparison(self):
        make_pdf(self.input, ["1234567890123456789012345678901234567890"])
        result = self.run_worker("施工资料.pdf")
        page = result["pages"][0]
        self.assertEqual(page["method"], "ocr")
        self.assertIn("suspected_missing_chinese", page["warnings"])
        self.assertIn("chinese_recovered_by_ocr", page["warnings"])
        self.assertIn("numeric_text_ocr_disagreement", page["warnings"])
        self.assertIn("[Extracted text]", self.report_text())
        self.assertIn("confidence=0.9800", self.report_text())

    def test_missing_chinese_not_resolved_by_more_ascii(self):
        make_pdf(self.input, [NORMAL_TEXT])
        self.ocr.return_value = [[[], NORMAL_TEXT * 2, 0.99]]
        result = self.run_worker("中文文件.pdf")
        self.assertIn("missing_chinese_unresolved_after_ocr_comparison", result["pages"][0]["warnings"])
        self.assertEqual(result["status"], "needs_review")

    def test_ink_without_text_is_not_blank(self):
        make_pdf(self.input, ["scan"])
        result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "ocr")
        self.assertTrue((self.output / "page-0001.png").is_file())
        self.assertEqual(result["ocrPages"], 1)
        self.ocr.assert_called_once()

    def test_real_mixed_image_page_requires_ocr(self):
        make_pdf(self.input, ["mixed"])
        result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "ocr")
        self.assertIn("mixed_image_page", result["pages"][0]["warnings"])

    def test_pdf_embedded_image_decode_pixels_are_bounded(self):
        make_pdf(self.input, ["mixed"])
        with patch.object(pypdfium2.PdfImage, "get_px_size", return_value=(100000, 100000)):
            result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "failed")
        self.assertIn("decoded pixel limit", " ".join(result["pages"][0]["warnings"]))
        self.ocr.assert_not_called()

    def test_pdf_text_error_falls_back_to_ocr_not_blank(self):
        make_pdf(self.input, [None])
        with patch.object(pypdfium2.PdfPage, "get_textpage", side_effect=RuntimeError("text decoder failed")):
            result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "ocr")
        self.assertTrue(any("text_extraction_failed" in w for w in result["pages"][0]["warnings"]))
        self.ocr.assert_called_once()

    def test_render_failure_is_not_blank(self):
        make_pdf(self.input, [None])
        with patch.object(pypdfium2.PdfPage, "render", side_effect=RuntimeError("render failed")):
            result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "failed")
        self.assertEqual(result["status"], "failed")

    def test_annotations_cannot_be_silently_blank(self):
        make_pdf(self.input, [None])
        self.ocr.return_value = []
        with patch.object(pypdfium2.raw, "FPDFPage_GetAnnotCount", return_value=1):
            result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "failed")
        self.ocr.assert_called_once()

    def test_pdf_rendered_pixels_are_bounded(self):
        writer = PdfWriter()
        writer.add_blank_page(width=100000, height=80000)
        writer.write(self.input)
        with patch.object(worker, "is_blank", return_value=False):
            result = self.run_worker()
        self.assertEqual(result["pages"][0]["method"], "ocr")
        image = self.ocr.call_args.args[0]
        self.assertLessEqual(image.width * image.height, worker.MAX_PIXELS)
        self.assertLessEqual(max(image.size), worker.MAX_DIMENSION)

    def test_missing_ocr_is_explicit_failure(self):
        make_pdf(self.input, ["scan"])
        with patch.dict(sys.modules, {"rapidocr_onnxruntime": None}):
            result = self.run_worker(ocr=worker.LocalOCR())
        self.assertEqual(result["status"], "needs_review")
        self.assertEqual(result["pages"][0]["method"], "failed")
        self.assertIn("Missing local OCR dependency", " ".join(result["pages"][0]["warnings"]))
        self.assertTrue((self.output / "page-0001.png").exists())

    def test_local_ocr_lazy_init_offline_models_and_thread_limits(self):
        models = self.root / "models"
        models.mkdir()
        for name in ("ch_PP-OCRv4_det_infer.onnx", "ch_ppocr_mobile_v2.0_cls_infer.onnx", "ch_PP-OCRv4_rec_infer.onnx"):
            (models / name).write_bytes(b"mock model")
        engine = Mock(return_value=([[[], "offline", 0.99]], None))
        constructor = Mock(return_value=engine)
        module = SimpleNamespace(__file__=str(self.root / "__init__.py"), RapidOCR=constructor)
        fake_cv = Mock()
        with patch.dict(sys.modules, {"rapidocr_onnxruntime": module, "cv2": fake_cv, "numpy": MagicMock()}), \
                patch("socket.create_connection", side_effect=AssertionError("network forbidden")):
            ocr = worker.LocalOCR()
            constructor.assert_not_called()
            for _ in range(2):
                self.assertEqual(ocr(Image.new("RGB", (100, 50))), [[[], "offline", 0.99]])
        constructor.assert_called_once()
        options = constructor.call_args.kwargs
        self.assertEqual(options["intra_op_num_threads"], 2)
        self.assertEqual(options["inter_op_num_threads"], 1)
        self.assertEqual(options["det_limit_type"], "max")
        self.assertEqual(options["text_score"], 0)
        self.assertFalse(options["det_use_cuda"])
        self.assertTrue(Path(options["det_model_path"]).is_absolute())
        fake_cv.setNumThreads.assert_called_once_with(2)

    def test_missing_bundled_model_never_downloads(self):
        constructor = Mock()
        module = SimpleNamespace(__file__=str(self.root / "__init__.py"), RapidOCR=constructor)
        with patch.dict(sys.modules, {"rapidocr_onnxruntime": module}), \
                patch("socket.create_connection", side_effect=AssertionError("network forbidden")):
            with self.assertRaisesRegex(worker.DiagnosticError, "Missing bundled local"):
                worker.LocalOCR()(Image.new("RGB", (100, 50)))
        constructor.assert_not_called()

    def test_low_confidence_never_ready(self):
        make_pdf(self.input, ["scan"])
        self.ocr.return_value = [[[], "123", 0.99], [[], "456", 0.35]]
        result = self.run_worker()
        page = result["pages"][0]
        self.assertEqual(page["confidence"], 0.35)
        self.assertIn("low_ocr_confidence", page["warnings"])
        self.assertIn("ocr_key_numbers_require_review", page["warnings"])
        self.assertEqual(result["status"], "needs_review")
        self.assertIn("confidence=0.3500", self.report_text())

    def test_empty_ocr_on_ink_is_needs_review_not_blank(self):
        make_pdf(self.input, ["scan"])
        self.ocr.return_value = []
        result = self.run_worker()
        self.assertEqual(result["status"], "needs_review")
        self.assertEqual(result["pages"][0]["method"], "failed")

    def test_nonfinite_ocr_confidence_is_needs_review(self):
        make_pdf(self.input, ["scan"])
        self.ocr.return_value = [[[], "bad", float("nan")]]
        self.assertEqual(self.run_worker()["status"], "needs_review")

    def test_partial_failure_is_review_and_processing_continues(self):
        make_pdf(self.input, ["scan", NORMAL_TEXT, None])
        self.ocr.side_effect = RuntimeError("offline model failure")
        events = []
        result = self.run_worker(emit=events.append)
        self.assertEqual(result["status"], "needs_review")
        self.assertEqual([e["page"] for e in events], [1, 2, 3])
        self.assertEqual([e["total"] for e in events], [3, 3, 3])
        self.assertEqual(result["pages"][0]["status"], "needs_review")
        self.assertEqual(result["pages"][1]["status"], "ready")

    def test_cache_reuses_ocr_and_rebuilds_report(self):
        make_pdf(self.input, ["scan", NORMAL_TEXT])
        self.run_worker()
        (self.output / "report.json").write_text("broken", encoding="utf-8")
        self.ocr.reset_mock()
        result = self.run_worker()
        self.assertEqual(result["cachedPages"], 2)
        self.ocr.assert_not_called()
        self.assertEqual(json.loads((self.output / "report.json").read_text(encoding="utf-8")), result)

    def test_failed_page_retries(self):
        make_pdf(self.input, ["scan"])
        self.ocr.side_effect = RuntimeError("temporarily missing engine")
        self.run_worker()
        self.ocr.side_effect = None
        result = self.run_worker()
        self.assertEqual(result["cachedPages"], 0)
        self.assertEqual(result["pages"][0]["method"], "ocr")
        self.assertEqual(self.ocr.call_count, 2)

    def test_corrupt_text_and_image_invalidate_cache(self):
        make_pdf(self.input, ["scan"])
        self.run_worker()
        for name in ("page-0001.txt", "page-0001.png"):
            with self.subTest(name=name):
                (self.output / name).write_bytes(b"corrupted")
                result = self.run_worker()
                self.assertEqual(result["cachedPages"], 0)
        self.assertEqual(self.ocr.call_count, 3)

    def test_policy_engine_dependency_and_name_invalidate_cache(self):
        make_pdf(self.input, ["scan"])
        self.run_worker()
        for field in ("policyVersion", "engineVersion", "dependencies", "sourceName", "sourceHash"):
            with self.subTest(field=field):
                path = self.output / "page-0001.json"
                cache = json.loads(path.read_text(encoding="utf-8"))
                cache["signature"][field] = "obsolete"
                path.write_text(json.dumps(cache), encoding="utf-8")
                self.assertEqual(self.run_worker()["cachedPages"], 0)

    def test_source_change_invalidates_cache(self):
        make_pdf(self.input, [NORMAL_TEXT])
        self.run_worker()
        make_pdf(self.input, ["scan"])
        self.assertEqual(self.run_worker()["cachedPages"], 0)

    def test_malformed_cache_is_not_trusted(self):
        make_pdf(self.input, ["scan"])
        self.run_worker()
        path = self.output / "page-0001.json"
        for value in ("[]", "null", "not json", '{"signature":{}}', "[" * 2000 + "0" + "]" * 2000):
            with self.subTest(value=value):
                path.write_text(value, encoding="utf-8")
                self.assertEqual(self.run_worker()["cachedPages"], 0)

    def test_cache_paths_and_status_are_validated(self):
        make_pdf(self.input, ["scan"])
        self.run_worker()
        path = self.output / "page-0001.json"
        for field, value in (("textPath", "../outside.txt"), ("imagePath", "C:\\secret.png"),
                             ("confidence", -1), ("status", "ready"), ("page", True), ("warnings", "bad"), ("warnings", [])):
            with self.subTest(field=field):
                cache = json.loads(path.read_text(encoding="utf-8"))
                cache["entry"][field] = value
                path.write_text(json.dumps(cache), encoding="utf-8")
                self.assertEqual(self.run_worker()["cachedPages"], 0)

    def test_hash_mismatch_fails_before_extraction(self):
        make_pdf(self.input, [NORMAL_TEXT])
        result = self.run_worker(source_hash="0" * 64)
        self.assertEqual(result["status"], "failed")
        self.assertIn("SHA-256 mismatch", " ".join(result["warnings"]))
        self.ocr.assert_not_called()

    def test_invalid_hash_rejected(self):
        make_pdf(self.input, [NORMAL_TEXT])
        with self.assertRaises(worker.DiagnosticError):
            self.run_worker(source_hash="../bad")

    def test_file_size_and_page_limits(self):
        make_pdf(self.input, [NORMAL_TEXT, NORMAL_TEXT])
        with patch.object(worker, "MAX_PAGES", 1):
            self.assertEqual(self.run_worker()["status"], "failed")
        with self.input.open("wb") as stream:
            stream.truncate(worker.MAX_FILE + 1)
        result = self.run_worker(source_hash="0" * 64)
        self.assertEqual(result["status"], "failed")
        self.assertIn("File exceeds size limit", " ".join(result["warnings"]))

    def test_corrupt_pdf_is_not_successful_blank(self):
        self.input.write_bytes(b"%PDF-broken")
        result = self.run_worker()
        self.assertEqual(result["status"], "failed")
        self.assertFalse(result["pages"])

    def test_docx_text_tables_headers_footers_without_ocr(self):
        source = self.root / "input.docx"
        make_docx(source)
        result = self.run_worker("文档.docx", source)
        self.assertEqual(result["status"], "ready")
        text = self.report_text()
        for expected in ("短中文正文", "项目\t42", "HEADER preserved", "FOOTER preserved", "not Word page"):
            self.assertIn(expected, text)
        self.assertIn("docx_page_numbers_are_logical_blocks_not_Word_pages", result["warnings"])
        self.ocr.assert_not_called()

    def test_docx_embedded_image_ocr_has_media_source(self):
        source = self.root / "input.docx"
        make_docx(source, picture=True)
        result = self.run_worker("test.docx", source)
        self.assertEqual(result["ocrPages"], 1)
        page = next(p for p in result["pages"] if p["method"] == "ocr")
        self.assertIn("media_has_no_physical_page_number", page["warnings"])
        self.assertIn("word/media/image1.png", " ".join(page["warnings"]))
        self.assertIn("word/media/image1.png", self.report_text())
        self.assertEqual(self.run_worker("test.docx", source)["cachedPages"], result["pageCount"])

    def test_docx_unsupported_image_is_failed_not_blank(self):
        source = self.root / "bad.docx"
        with zipfile.ZipFile(source, "w") as z:
            z.writestr("word/document.xml", f'<w:document xmlns:w="{worker.W[1:-1]}"><w:body/></w:document>')
            z.writestr("word/media/image1.emf", b"not a supported raster")
        result = self.run_worker("bad.docx", source)
        self.assertEqual(result["status"], "failed")
        self.assertEqual(result["pages"][0]["method"], "failed")

    def test_docx_multiframe_media_is_not_a_blank_first_frame(self):
        source = self.root / "frames.docx"
        first = Image.new("RGB", (100, 50), "white")
        second = Image.new("RGB", (100, 50), "black")
        media = io.BytesIO()
        first.save(media, "GIF", save_all=True, append_images=[second])
        with zipfile.ZipFile(source, "w") as z:
            z.writestr("word/document.xml", f'<w:document xmlns:w="{worker.W[1:-1]}"><w:body/></w:document>')
            z.writestr("word/media/frames.gif", media.getvalue())
        result = self.run_worker("frames.docx", source)
        self.assertEqual(result["pages"][0]["method"], "failed")
        self.assertIn("Multi-frame", " ".join(result["pages"][0]["warnings"]))

    def test_docx_eps_cannot_invoke_postscript_interpreter(self):
        source = self.root / "postscript.docx"
        with zipfile.ZipFile(source, "w") as z:
            z.writestr("word/document.xml", f'<w:document xmlns:w="{worker.W[1:-1]}"><w:body/></w:document>')
            z.writestr("word/media/image.eps", b"%!PS-Adobe-3.0 EPSF-3.0\n%%BoundingBox: 0 0 100 100\nshowpage\n")
        with patch("subprocess.Popen", side_effect=AssertionError("external interpreter forbidden")) as spawn:
            result = self.run_worker("postscript.docx", source)
        self.assertEqual(result["pages"][0]["method"], "failed")
        spawn.assert_not_called()

    def test_docx_missing_media_or_unsupported_content_is_not_blank(self):
        for content in ("<w:drawing/>", "<w:object/>", "<w:altChunk/>"):
            with self.subTest(content=content):
                source = self.root / "unsupported.docx"
                with zipfile.ZipFile(source, "w") as z:
                    z.writestr("word/document.xml", f'<w:document xmlns:w="{worker.W[1:-1]}"><w:body><w:p>{content}</w:p></w:body></w:document>')
                result = self.run_worker("unsupported.docx", source)
                self.assertEqual(result["pages"][0]["method"], "failed")

    def test_docx_missing_media_with_normal_text_requires_review(self):
        source = self.root / "missing-media.docx"
        make_docx(source, picture=True)
        with zipfile.ZipFile(source) as archive:
            entries = {name: archive.read(name) for name in archive.namelist() if not name.startswith("word/media/")}
        with zipfile.ZipFile(source, "w") as archive:
            for name, data in entries.items():
                archive.writestr(name, data)
        result = self.run_worker("missing-media.docx", source)
        self.assertEqual(result["status"], "needs_review")
        self.assertTrue(any(p["method"] == "failed" for p in result["pages"]))
        self.assertIn("短中文正文", self.report_text())
        self.ocr.assert_not_called()

    def test_docx_footer_missing_image_relationship_requires_review(self):
        source = self.root / "footer-missing-image.docx"
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("word/document.xml",
                             f'<w:document xmlns:w="{worker.W[1:-1]}"><w:body><w:p><w:r><w:t>正文完整</w:t></w:r></w:p></w:body></w:document>')
            archive.writestr("word/footer1.xml",
                             f'<w:ftr xmlns:w="{worker.W[1:-1]}" xmlns:r="{worker.R[1:-1]}" '
                             'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                             '<w:p><w:r><w:t>页脚文字</w:t></w:r></w:p>'
                             '<w:p><w:r><w:drawing><a:blip r:embed="rIdMissing"/></w:drawing></w:r></w:p>'
                             '</w:ftr>')
        result = self.run_worker("页脚缺图.docx", source)
        missing = next(page for page in result["pages"]
                       if "missing_embedded_image_relationship" in page["warnings"])
        self.assertEqual(result["status"], "needs_review")
        self.assertEqual(missing["method"], "blank")
        self.assertEqual(missing["status"], "needs_review")
        self.assertIn("embedded_media_source: word/footer1.xml", missing["warnings"])
        self.assertFalse(any(page["status"] == "failed" for page in result["pages"]))
        self.assertIn("正文完整", self.report_text())
        self.assertEqual(self.run_worker("页脚缺图.docx", source)["cachedPages"], result["pageCount"])

    def test_empty_docx_is_blank_logical_block(self):
        source = self.root / "empty.docx"
        Document().save(source)
        result = self.run_worker("empty.docx", source)
        self.assertEqual(result["pages"][0]["method"], "blank")
        self.assertIn("not a Word page", self.report_text())

    def test_docx_dtd_and_entities_rejected_including_utf16(self):
        payload = '<!DOCTYPE document [<!ENTITY x "unsafe">]><document>&x;</document>'
        for encoding in ("utf-8", "utf-16"):
            with self.subTest(encoding=encoding):
                source = self.root / "entities.docx"
                with zipfile.ZipFile(source, "w") as z:
                    z.writestr("word/document.xml", payload.encode(encoding))
                result = self.run_worker("entities.docx", source)
                self.assertEqual(result["status"], "failed")
                self.assertIn("DTD/entities", " ".join(result["warnings"]))

    def test_zip_path_entries_and_limits_rejected(self):
        source = self.root / "bad.docx"
        with zipfile.ZipFile(source, "w") as z:
            z.writestr("../escape.xml", "<test/>")
        result = self.run_worker("bad.docx", source)
        self.assertEqual(result["status"], "failed")
        make_docx(source)
        with patch.object(worker, "MAX_ZIP_ENTRIES", 1):
            self.assertEqual(self.run_worker("bad.docx", source)["status"], "failed")
        with patch.object(worker, "MAX_UNPACKED", 1):
            self.assertEqual(self.run_worker("bad.docx", source)["status"], "failed")

    def test_external_docx_links_are_not_followed(self):
        source = self.root / "external.docx"
        with zipfile.ZipFile(source, "w") as z:
            z.writestr("word/document.xml", f'<w:document xmlns:w="{worker.W[1:-1]}"><w:body><w:p><w:r><w:t>Safe text</w:t></w:r></w:p></w:body></w:document>')
            z.writestr("word/_rels/document.xml.rels", '<Relationships><Relationship TargetMode="External" Target="file:///does-not-exist"/></Relationships>')
            z.writestr("word/vbaProject.bin", b"never execute")
        result = self.run_worker("external.docx", source)
        self.assertIn("external_links_not_followed", result["warnings"])
        self.assertTrue(any("active_or_embedded_content_ignored" in w for w in result["warnings"]))
        self.assertIn("Safe text", self.report_text())

    def test_blank_detection_and_image_size_bounds(self):
        white = Image.new("RGB", (500, 500), "white")
        self.assertTrue(worker.is_blank(white))
        white.putpixel((10, 10), (0, 0, 0))
        self.assertFalse(worker.is_blank(white))
        large = Image.new("RGB", (5000, 3000), "white")
        reduced = worker.bounded_image(large)
        self.assertLessEqual(reduced.width * reduced.height, worker.MAX_PIXELS)
        self.assertLessEqual(max(reduced.size), worker.MAX_DIMENSION)
        with patch.object(worker, "MAX_IMAGE_PIXELS", 100):
            with self.assertRaises(worker.DiagnosticError):
                worker.bounded_image(white)

    def test_quality_heuristics(self):
        self.assertEqual(worker.text_issues(NORMAL_TEXT, "English.pdf"), [])
        self.assertIn("sparse_text", worker.text_issues("short", "English.pdf"))
        self.assertIn("garbled_text", worker.text_issues(NORMAL_TEXT + "\ufffd", "English.pdf"))
        self.assertIn("mixed_image_page", worker.text_issues(NORMAL_TEXT, "English.pdf", True))

    def make_link(self, link, target, directory=False):
        try:
            link.symlink_to(target, target_is_directory=directory)
        except OSError as exc:
            self.skipTest("Symlink creation not permitted by this host: " + str(exc))

    def test_symlink_output_directory_is_rejected(self):
        make_pdf(self.input, [NORMAL_TEXT])
        outside = self.root / "outside"
        outside.mkdir()
        self.make_link(self.output, outside, True)
        with self.assertRaises((worker.UnsafePath, OSError)):
            self.run_worker()
        self.assertFalse(list(outside.iterdir()))

    def test_symlink_output_artifacts_are_rejected(self):
        make_pdf(self.input, ["scan"])
        self.output.mkdir()
        target = self.root / "untouched.txt"
        target.write_bytes(b"untouched")
        for name in ("report.json", "full.txt", "page-0001.json", "page-0001.txt", "page-0001.png"):
            with self.subTest(name=name):
                link = self.output / name
                self.make_link(link, target)
                try:
                    with self.assertRaises((worker.UnsafePath, OSError)):
                        self.run_worker()
                    self.assertEqual(target.read_bytes(), b"untouched")
                finally:
                    link.unlink()

    def test_symlink_input_is_rejected(self):
        source = self.root / "real.pdf"
        make_pdf(source, [NORMAL_TEXT])
        self.make_link(self.input, source)
        with self.assertRaises((worker.UnsafePath, OSError)):
            self.run_worker()

    def test_hardlink_output_is_rejected(self):
        make_pdf(self.input, [NORMAL_TEXT])
        self.output.mkdir()
        target = self.root / "untouched.txt"
        target.write_bytes(b"untouched")
        os.link(target, self.output / "full.txt")
        with self.assertRaises(worker.UnsafePath):
            self.run_worker()
        self.assertEqual(target.read_bytes(), b"untouched")

    def test_atomic_replace_leaves_old_file_on_failure(self):
        self.output.mkdir()
        target = self.output / "full.txt"
        target.write_bytes(b"old")
        with worker.SafeDirectory(self.output) as output:
            with patch.object(worker.os, "replace", side_effect=OSError("simulated failure")):
                with self.assertRaises(OSError):
                    output.write("full.txt", b"new")
        self.assertEqual(target.read_bytes(), b"old")
        self.assertEqual([p.name for p in self.output.iterdir()], ["full.txt"])

    def test_relative_and_traversal_paths_are_rejected(self):
        for path in ("relative", str(self.root / ".." / "escape")):
            with self.subTest(path=path), self.assertRaises(worker.UnsafePath):
                worker.SafeDirectory(path, create=True)
        self.output.mkdir()
        with worker.SafeDirectory(self.output) as output:
            for name in ("../escape", "x:y", "sub/file", ".."):
                with self.subTest(name=name), self.assertRaises(worker.UnsafePath):
                    output.write(name, b"bad")

    def test_input_cannot_be_overwritten_by_output(self):
        self.output.mkdir()
        source = self.output / "full.txt"
        make_pdf(source, [NORMAL_TEXT])
        original = source.read_bytes()
        with self.assertRaises(worker.UnsafePath):
            self.run_worker(source=source)
        self.assertEqual(source.read_bytes(), original)

    def test_cli_isolated_jsonl_and_no_input_directory_imports(self):
        make_pdf(self.input, [NORMAL_TEXT, None])
        (self.root / "pypdfium2.py").write_text('raise RuntimeError("untrusted import")', encoding="utf-8")
        result = subprocess.run([sys.executable, "-I", "-B", str(WORKER_PATH), "--input", str(self.input),
                                 "--output", str(self.output), "--sha256", hashlib.sha256(self.input.read_bytes()).hexdigest(),
                                 "--name", "example.pdf"], cwd=self.root, capture_output=True, text=True, timeout=60)
        self.assertEqual(result.returncode, 0, result.stderr)
        events = [json.loads(line) for line in result.stdout.splitlines()]
        self.assertEqual(events[-1], {"type": "done"})
        self.assertEqual([event["page"] for event in events[:-1]], [1, 2])

    def test_cli_failure_has_done_diagnostic_and_nonzero_exit(self):
        make_pdf(self.input, [NORMAL_TEXT])
        result = subprocess.run([sys.executable, "-I", "-B", str(WORKER_PATH), "--input", str(self.input),
                                 "--output", str(self.output), "--sha256", "0" * 64,
                                 "--name", "example.pdf"], capture_output=True, text=True, timeout=60)
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(json.loads(result.stdout.strip()), {"type": "done"})
        self.assertIn("SHA-256 mismatch", result.stderr)

    def test_protocol_progress_is_flushed(self):
        make_pdf(self.input, [NORMAL_TEXT])
        stream = io.StringIO()
        with patch.object(worker.sys, "stdout", stream), patch.object(worker, "print", wraps=print) as printing:
            code = worker.main(["--input", str(self.input), "--output", str(self.output),
                                "--sha256", hashlib.sha256(self.input.read_bytes()).hexdigest(), "--name", "example.pdf"])
        self.assertEqual(code, 0)
        self.assertTrue(all(call.kwargs.get("flush") for call in printing.call_args_list))
        self.assertEqual(json.loads(stream.getvalue().splitlines()[-1]), {"type": "done"})


if __name__ == "__main__":
    unittest.main(verbosity=2)
